"""Create a copyable DOCX relational schema from the canonical DDL."""
from pathlib import Path
import re, json, hashlib
from zipfile import ZipFile
import xml.etree.ElementTree as ET
from docx import Document
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_UNDERLINE, WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2];OUT=Path(__file__).resolve().parent
DDL=ROOT/'database/dbms12-compact/01_schema.sql'
sql=DDL.read_text(encoding='utf-8-sig')
fks={(a,b):(c,d) for a,b,c,d in re.findall(r'ALTER TABLE dbo\.\[(\w+)\].*?FOREIGN KEY \(\[(\w+)\]\) REFERENCES dbo\.\[(\w+)\] \(\[(\w+)\]\)',sql)}
tables=[]
for name,body in re.findall(r'CREATE TABLE dbo\.\[(\w+)\] \((.*?)\n\);',sql,re.S):
    cols=[m[1] for line in body.splitlines() if (m:=re.match(r'\s*\[(\w+)\]\s+(?:int|bigint|tinyint|varchar|nvarchar|date|datetime2|decimal|binary|varbinary|bit|rowversion|uniqueidentifier)\b',line))]
    pk=re.findall(r'\[(\w+)\]',re.search(r'PRIMARY KEY \((.*?)\)',body)[1])
    tables.append((name,cols,pk))
assert len(tables)==12 and sum(len(t[1]) for t in tables)==99 and len(fks)==19
assert sum(len(t[2]) for t in tables)==12
assert not any(c in pk and (name,c) in fks for name,cols,pk in tables for c in cols)

doc=Document();sec=doc.sections[0]
sec.page_width=Mm(210);sec.page_height=Mm(297)
sec.top_margin=sec.bottom_margin=Mm(18)
sec.left_margin=sec.right_margin=Mm(18)
for name in ['Normal','Title','TitleChar']:
    st=doc.styles[name];st.font.name='Times New Roman';st.font.color.rgb=RGBColor(0,0,0)
    st.element.rPr.rFonts.set(qn('w:eastAsia'),'Times New Roman')
    fonts=st.element.rPr.rFonts
    for attr in list(fonts.attrib):
        if 'theme' in attr.lower():del fonts.attrib[attr]
    fonts.set(qn('w:cs'),'Times New Roman')
    for border in list(st.element.findall('.//'+qn('w:pBdr'))):border.getparent().remove(border)
normal=doc.styles['Normal'];normal.font.size=Pt(12)
normal.paragraph_format.space_before=Pt(0);normal.paragraph_format.space_after=Pt(5)
normal.paragraph_format.line_spacing=1.08
title=doc.add_paragraph('Lược đồ quan hệ DBMS12',style='Title')
title.alignment=WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_after=Pt(7)
for run in title.runs:run.bold=True;run.font.size=Pt(16);run.font.name='Times New Roman'
doc.add_paragraph('Lược đồ gồm 12 bảng và 99 thuộc tính. Khóa chính gạch chân liền; khóa ngoại gạch chân nét đứt.')
for name,cols,pk in tables:
    p=doc.add_paragraph();p.paragraph_format.keep_together=True
    p.add_run(name).bold=True;p.add_run('(')
    for i,col in enumerate(cols):
        if i:p.add_run(', ')
        run=p.add_run(col)
        if col in pk:run.underline=WD_UNDERLINE.SINGLE
        elif (name,col) in fks:run.underline=WD_UNDERLINE.DASH
        # Explicit language/font avoids theme and proofing residue in copied text.
        run.font.name='Times New Roman'
        rpr=run._r.get_or_add_rPr();lang=OxmlElement('w:lang');lang.set(qn('w:val'),'en-US');rpr.append(lang)
    p.add_run(').')
doc.core_properties.title='Lược đồ quan hệ DBMS12'
doc.core_properties.subject='12 bảng và 99 thuộc tính với ký hiệu khóa chính và khóa ngoại'
doc.core_properties.author='';doc.core_properties.last_modified_by=''
path=OUT/'Luoc-do-quan-he-DBMS12.docx';doc.save(path)

# Verify actual Word runs, including the dashed underline property, against DDL.
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with ZipFile(path) as z:xml=ET.fromstring(z.read('word/document.xml'))
paras=xml.findall('w:body/w:p',ns)[2:]
assert len(paras)==12
counts={'single':0,'dash':0,'ordinary':0}
for (name,cols,pk),p in zip(tables,paras):
    text=''.join(t.text or '' for t in p.findall('.//w:t',ns))
    assert text==name+'('+', '.join(cols)+').',name
    attr_runs=[r for r in p.findall('w:r',ns) if (r.find('w:t',ns) is not None and r.find('w:t',ns).text in cols)]
    assert len(attr_runs)==len(cols)
    for col,r in zip(cols,attr_runs):
        u=r.find('w:rPr/w:u',ns);kind=u.get(qn('w:val')) if u is not None else 'ordinary'
        expected='single' if col in pk else 'dash' if (name,col) in fks else 'ordinary'
        assert kind==expected,(name,col,kind,expected)
        counts[kind]+=1
assert counts=={'single':12,'dash':19,'ordinary':68}
receipt={'date':'2026-10-22','tables':12,'columns':99,'wordUnderlineCounts':counts,
         'schemaSha256':hashlib.sha256(DDL.read_bytes()).hexdigest().upper(),
         'docxSha256':hashlib.sha256(path.read_bytes()).hexdigest().upper(),
         'render':'pending','format':'A4 portrait Times New Roman 12 pt; editable paragraphs'}
(ROOT/'docs/evidence/relational-word.json').write_text(json.dumps(receipt,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
print(json.dumps(receipt,ensure_ascii=False))
